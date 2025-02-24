"""
Stage 4 Excel Formatter

This script takes the stage 3 output CSV and creates a formatted Excel file with
color-coded cells based on values.

Input:
- output/stage_3_output.csv: Final data with all analysis

Output:
- output/final_report.xlsx: Formatted Excel file with color-coded cells
"""

import pandas as pd
import openpyxl
from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from openpyxl.styles import PatternFill, Font
from openpyxl.formatting.rule import CellIsRule
import logging
import os

# Configure logging
logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)

# Color definitions
COLORS = {
    'green': '90EE90',   # Light green
    'red': 'FFB6B6',     # Light red
    'yellow': 'FFE5B4',  # Light yellow
    'orange': 'FFB347'   # Light orange
}

def apply_stage_headers(worksheet):
    """Add stage labels above column groups."""
    # Define which columns belong to which stage
    stage_columns = {
        'Stage 1': [ 'Is Site Live', 'Keywords Found', 'Proof'],
        'Stage 2': ['Is Closed', 'Close Date', 'Close Reason'],
        'Stage 3': ['Latest Status', 'Engagement Score', 'Sentiment Score', 'Interest Level', 
                   'Key Objections', 'Latest Activity', 'Data Confidence', 'Positive Patterns', 
                   'Negative Patterns']
    }
    
    # Insert a new row at the top
    worksheet.insert_rows(1)
    
    # Get all column headers
    headers = [cell.value for cell in worksheet[2]]
    
    # Add stage labels
    for stage, columns in stage_columns.items():
        start_idx = None
        end_idx = None
        
        # Find the start and end columns for this stage
        for col_name in columns:
            if col_name in headers:
                idx = headers.index(col_name) + 1  # +1 because Excel is 1-based
                if start_idx is None:
                    start_idx = idx
                end_idx = idx
        
        if start_idx and end_idx:
            # Add stage label
            start_cell = worksheet.cell(row=1, column=start_idx)
            start_cell.value = stage
            
            # Style the stage label
            stage_fill = PatternFill(start_color='D3D3D3', end_color='D3D3D3', fill_type='solid')
            stage_font = Font(bold=True)
            
            # Apply style to the entire stage header
            for col in range(start_idx, end_idx + 1):
                cell = worksheet.cell(row=1, column=col)
                cell.fill = stage_fill
                cell.font = stage_font
            
            # Merge cells if stage spans multiple columns
            if end_idx > start_idx:
                worksheet.merge_cells(
                    start_row=1, 
                    start_column=start_idx, 
                    end_row=1, 
                    end_column=end_idx
                )

def apply_conditional_formatting(worksheet, start_row=3):
    """Apply conditional formatting rules to the worksheet."""
    
    # Get the header row to find column indices
    headers = [cell.value for cell in worksheet[2]]  # Headers are now in row 2
    
    # Binary columns (Yes/No, True/False)
    binary_columns = {
        'Is Site Live': {'Yes': 'green', 'No': 'red'},
        'Is Closed': {'True': 'red', 'False': 'green'}
    }
    
    # Interest Level formatting with orange for Medium
    interest_colors = {
        'High': 'green',
        'Medium': 'orange',
        'Low': 'yellow',
        'None': 'red'
    }
    
    # Data Confidence formatting
    confidence_colors = {
        'High': 'green',
        'Medium': 'yellow',
        'Low': 'red'
    }
    
    # Engagement Score formatting (numeric ranges)
    engagement_rules = [
        ('>=75', 'green'),
        ('>=25', 'yellow'),
        ('>0', 'red')
    ]
    
    # Sentiment Score formatting (numeric ranges)
    sentiment_rules = [
        ('>=0.5', 'green'),
        ('>=0', 'yellow'),
        ('<0', 'red')
    ]
    
    # Apply formatting for each column
    for col_idx, header in enumerate(headers, start=1):
        # Get column letter from index directly
        col_letter = get_column_letter(col_idx)
        
        # Binary columns
        if header in binary_columns:
            for value, color in binary_columns[header].items():
                worksheet.conditional_formatting.add(
                    f'{col_letter}{start_row}:{col_letter}{worksheet.max_row}',
                    CellIsRule(operator='equal', formula=[f'"{value}"'], fill=PatternFill(start_color=COLORS[color], end_color=COLORS[color], fill_type='solid'))
                )
        
        # Interest Level
        elif header == 'Interest Level':
            for level, color in interest_colors.items():
                worksheet.conditional_formatting.add(
                    f'{col_letter}{start_row}:{col_letter}{worksheet.max_row}',
                    CellIsRule(operator='equal', formula=[f'"{level}"'], fill=PatternFill(start_color=COLORS[color], end_color=COLORS[color], fill_type='solid'))
                )
        
        # Data Confidence
        elif header == 'Data Confidence':
            for level, color in confidence_colors.items():
                worksheet.conditional_formatting.add(
                    f'{col_letter}{start_row}:{col_letter}{worksheet.max_row}',
                    CellIsRule(operator='equal', formula=[f'"{level}"'], fill=PatternFill(start_color=COLORS[color], end_color=COLORS[color], fill_type='solid'))
                )
        
        # Engagement Score
        elif header == 'Engagement Score':
            for threshold, color in engagement_rules:
                worksheet.conditional_formatting.add(
                    f'{col_letter}{start_row}:{col_letter}{worksheet.max_row}',
                    CellIsRule(operator=threshold[0], formula=[threshold[1:]], fill=PatternFill(start_color=COLORS[color], end_color=COLORS[color], fill_type='solid'))
                )
        
        # Sentiment Score
        elif header == 'Sentiment Score':
            for threshold, color in sentiment_rules:
                worksheet.conditional_formatting.add(
                    f'{col_letter}{start_row}:{col_letter}{worksheet.max_row}',
                    CellIsRule(operator=threshold[0], formula=[threshold[1:]], fill=PatternFill(start_color=COLORS[color], end_color=COLORS[color], fill_type='solid'))
                )

def format_excel(input_path: str, output_path: str):
    """
    Read CSV and create a formatted Excel file.
    """
    try:
        # Read the CSV
        df = pd.read_csv(input_path)
        
        # Create Excel writer
        logger.info("Creating Excel file...")
        df.to_excel(output_path, index=False, engine='openpyxl')
        
        # Load the workbook for formatting
        workbook = openpyxl.load_workbook(output_path)
        worksheet = workbook.active
        
        # Add stage headers
        logger.info("Adding stage headers...")
        apply_stage_headers(worksheet)
        
        # Format column header row (now row 2)
        header_fill = PatternFill(start_color='4F81BD', end_color='4F81BD', fill_type='solid')
        header_font = Font(color='FFFFFF', bold=True)
        
        for cell in worksheet[2]:
            cell.fill = header_fill
            cell.font = header_font
        
        # Apply conditional formatting
        logger.info("Applying conditional formatting...")
        apply_conditional_formatting(worksheet)
        
        # Auto-adjust column widths
        logger.info("Adjusting column widths...")
        for col_idx in range(1, worksheet.max_column + 1):
            max_length = 0
            column_letter = get_column_letter(col_idx)
            
            # Check all cells in column
            for row in range(1, worksheet.max_row + 1):
                cell = worksheet.cell(row=row, column=col_idx)
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            
            adjusted_width = min(max_length + 2, 50)  # Cap width at 50
            worksheet.column_dimensions[column_letter].width = adjusted_width
        
        # Freeze the header rows
        worksheet.freeze_panes = 'A3'
        
        # Save the formatted workbook
        workbook.save(output_path)
        logger.info(f"Formatted Excel file saved to: {output_path}")
        
    except Exception as e:
        logger.error(f"Error formatting Excel file: {str(e)}")
        raise

def main():
    """Main execution function."""
    try:
        logger.info("\nStarting stage 4 processing...")
        
        input_path = 'output/stage_3_output.csv'
        output_path = 'output/final_report.xlsx'
        
        # Ensure output directory exists
        os.makedirs(os.path.dirname(output_path), exist_ok=True)
        
        # Validate input file exists
        if not os.path.exists(input_path):
            raise FileNotFoundError(f"Stage 3 output not found at: {input_path}")
        
        # Format the Excel file
        format_excel(input_path, output_path)
        
        # Verify output
        if os.path.exists(output_path):
            file_size = os.path.getsize(output_path)
            logger.info(f"\nOutput file size: {file_size} bytes")
        else:
            logger.warning("Warning: Output file was not created")
        
    except Exception as e:
        logger.error(f"An error occurred: {str(e)}")
        raise

if __name__ == "__main__":
    main() 